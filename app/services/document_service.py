from pathlib import Path
from uuid import uuid4

from ..config import settings
from ..database.database import db
from ..errors import AppError

ALLOWED = {".pdf", ".txt", ".md", ".markdown", ".docx"}


def safe_name(filename):
    name = Path((filename or "").replace("\\", "/")).name.strip()  # drops any directory part
    ext = Path(name).suffix.lower()
    if not name or ext not in ALLOWED:
        raise AppError("unsupported_file", "Unsupported file type. Upload PDF, TXT, DOCX or Markdown files.", 415)
    return name[:200], ext


def save_file(data: bytes, ext: str) -> Path:
    settings.docs_dir.mkdir(parents=True, exist_ok=True)
    p = settings.docs_dir / f"{uuid4().hex}{ext}"  # server-generated name: no path tricks; never executed
    p.write_bytes(data)
    return p


def extract_text(path: Path, ext: str) -> str:
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            text = "\n\n".join((pg.extract_text() or "") for pg in PdfReader(str(path)).pages)
        elif ext == ".docx":
            import docx
            d = docx.Document(str(path))
            text = "\n".join(p.text for p in d.paragraphs)
            for t in d.tables:
                for row in t.rows:
                    text += "\n" + " | ".join(c.text for c in row.cells)
        else:
            text = path.read_text(encoding="utf-8", errors="replace")
    except Exception:
        raise AppError("parse_failed",
                       "Could not read this file. It may be corrupted, password-protected or not a real document.", 422)
    if not text.strip():
        raise AppError("empty_document",
                       "No readable text found. Scanned/image-only PDFs need OCR, which is not included.", 422)
    return text


def list_documents() -> list[dict]:
    with db() as c:
        rows = c.execute("SELECT id,name,ext,size,chunk_count,created_at FROM documents ORDER BY id DESC").fetchall()
    return [dict(r) for r in rows]


def delete_document(doc_id: int) -> bool:
    with db() as c:
        r = c.execute("SELECT stored_name FROM documents WHERE id=?", (doc_id,)).fetchone()
        if not r:
            return False
        c.execute("DELETE FROM documents WHERE id=?", (doc_id,))
    (settings.docs_dir / r["stored_name"]).unlink(missing_ok=True)
    return True
